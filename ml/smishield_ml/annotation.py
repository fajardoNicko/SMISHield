"""Excel workbooks for manual labelling (see ml/LABELING_GUIDE.md)."""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

LABEL_OPTIONS = ["Safe", "Suspicious", "High Risk", "Exclude"]

INSTRUCTIONS = [
    "Label each message using ml/LABELING_GUIDE.md. Work alone; do not discuss labels until the agreement check is done.",
    "",
    "label     Pick Safe, Suspicious, High Risk, or Exclude from the drop-down.",
    "criteria  Write the guide's codes that apply, e.g. H1; H2  or  S1; S2  or  F1.",
    "notes     Anything unclear (optional).",
    "",
    "Decide in this order and stop at the first that applies:",
    "  0. Exclude: empty, unreadable, not Filipino/English/Taglish, or not an SMS.",
    "  1. High Risk if ANY of:",
    "     H1 asks for sensitive information (OTP, PIN, password, card/account details, ID, selfie) or to enter it via a link",
    "     H2 poses as a bank, e-wallet, agency, courier, telco or a person you know AND asks you to click, call, reply or pay",
    "     H3 asks for money, a fee, load or gift cards under a pretext (prize, parcel, loan, penalty, emergency)",
    "     H4 makes a false claim about your account or money (on hold, money waiting for you) to make you act",
    "  2. Suspicious if ANY of (and no H criterion):",
    "     S1 gambling or betting promotion     S2 too-good-to-be-true offer: prize, free money, bonus",
    "     S3 moves you to Telegram/WhatsApp or an unknown contact     S4 unknown or shortened link",
    "     S5 loan, investment, crypto or easy-job offer     S6 urgency or time pressure",
    "  3. Safe otherwise: F1 genuine transaction/OTP/notification, F2 official advisory, F3 provider's own promo,",
    "     F4 personal or everyday message, F5 unwanted but harmless (e.g. political campaign text with no H/S criterion).",
    "",
    "Use the sender column only to judge whether a message really comes from the company it names.",
    "If the content meets an H criterion, label High Risk even when the sender looks official (sender IDs can be spoofed).",
]

_HEADER_FILL = PatternFill("solid", fgColor="1F3A68")
_WIDTHS = {"id": 11, "sender": 22, "message": 90, "label": 13, "criteria": 14,
           "notes": 30, "annotator": 11}


def sender_display(sender_type, sender_name):
    if sender_type in ("alphanumeric", "shortcode") and sender_name:
        kind = "sender ID" if sender_type == "alphanumeric" else "short code"
        return f"{sender_name} ({kind})"
    return {"mobile": "mobile number", "email": "e-mail address"}.get(sender_type, "unknown")


def has_labels(path):
    path = Path(path)
    if not path.exists():
        return False
    df = read_sheet(path)
    label_cols = [c for c in df.columns if c == "label" or c == "final_label"]
    return any(df[c].notna().any() for c in label_cols)


def write_sheet(path, rows, columns, label_column="label", title="Messages"):
    """Writes rows (DataFrame) to an annotation workbook with a label drop-down."""
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(columns)
    for record in rows[columns].itertuples(index=False):
        ws.append(["" if pd.isna(v) else v for v in record])

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
    for i, name in enumerate(columns, start=1):
        letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[letter].width = _WIDTHS.get(name, 14)
        if name == "message":
            for cell in ws[letter][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"

    if label_column in columns:
        letter = ws.cell(row=1, column=columns.index(label_column) + 1).column_letter
        rule = DataValidation(type="list", formula1='"' + ",".join(LABEL_OPTIONS) + '"',
                              allow_blank=True, showErrorMessage=True,
                              errorTitle="Invalid label", error="Choose a label from the list.")
        ws.add_data_validation(rule)
        rule.add(f"{letter}2:{letter}{len(rows) + 1}")

    guide = wb.create_sheet("Instructions")
    for line in INSTRUCTIONS:
        guide.append([line])
    guide.column_dimensions["A"].width = 140

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def read_sheet(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return pd.DataFrame()
    df = pd.DataFrame(rows[1:], columns=rows[0])
    for col in ("label", "final_label"):
        if col in df.columns:
            df[col] = df[col].map(lambda v: str(v).strip() if v is not None and str(v).strip() else None)
    return df
